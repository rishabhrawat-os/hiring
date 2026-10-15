#!/usr/bin/env python3
"""Render a jd.json into a .docx in the house layout (brand/ + company/).

Layout, fonts and colours come from brand/tokens.json; fixed text from
structure/skeleton.json and company/company.json. The three brand fonts are
embedded in the file so it looks the same on a laptop without them.

    python3 render_docx.py jd.json out.docx
"""
import json
import re
import sys
import uuid
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_LINE_SPACING
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement, parse_xml
    from lxml import etree
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing: pip install python-docx")

ROOT = Path(__file__).resolve().parent.parent
TOK = json.loads((ROOT / "brand/tokens.json").read_text())
SKEL = json.loads((ROOT / "structure/skeleton.json").read_text())
CO = json.loads((ROOT / "company/company.json").read_text())
C = TOK["color"]
F = TOK["font"]
T = TOK["text"]
P = TOK["page"]


# ---------- low-level helpers ----------

# Child order required by the OOXML schema. Word rejects files that break it.
ORDER = {
    "w:pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
             "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
             "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
             "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "w:rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
             "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
             "em lang eastAsianLayout specVanish oMath",
    "w:tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "w:tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
               "tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "w:settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
                  "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData "
                  "embedTrueTypeFonts embedSystemFonts saveSubsetFonts",
}


def put(parent, child):
    """Insert child into parent at its schema position, replacing any element with the same tag."""
    pt = parent.tag.split("}")[1]
    order = ORDER["w:" + pt].split()
    name = child.tag.split("}")[1]
    for old in parent.findall(child.tag):
        parent.remove(old)
    mine = order.index(name)
    for i, c in enumerate(parent):
        cname = c.tag.split("}")[1] if "}" in c.tag else c.tag
        if cname not in order or order.index(cname) > mine:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def tw(pt):
    """points -> twentieths of a point (twips)"""
    return int(round(pt * 20))


def style_run(run, spec, **over):
    s = dict(spec, **over)
    font = F[s.get("font", "serif")]
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = put(rpr, el("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    run.font.size = Pt(s["size"])
    run.font.color.rgb = RGBColor.from_string(C[s.get("color", "ink")])
    run.font.bold = bool(s.get("bold"))
    run.font.italic = bool(s.get("italic"))
    if s.get("tracking"):
        put(rpr, el("w:spacing", **{"w:val": tw(s["tracking"])}))
    if s.get("case") == "upper":
        run.text = run.text.upper()
    return run


def para(container, spec=None, before=0, after=0, lh=None, indent=0, hanging=0, right=0, keep=False, align=None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    line = lh if lh is not None else (spec or {}).get("lh")
    if line:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    else:
        pf.line_spacing = 1.0
    if indent or hanging:
        pf.left_indent = Pt(indent)
        pf.first_line_indent = Pt(-hanging) if hanging else None
    if right:
        pf.right_indent = Pt(right)
    if keep:
        pf.keep_with_next = True
    pf.widow_control = True
    if align:
        pf.alignment = align
    return p


def border(p_or_cell_pr, side, color, size_pt, space=0):
    """Add a border to a paragraph (pPr) or table cell (tcPr)."""
    tag = "w:pBdr" if p_or_cell_pr.tag == qn("w:pPr") else "w:tcBorders"
    box = p_or_cell_pr.find(qn(tag))
    if box is None:
        box = put(p_or_cell_pr, el(tag))
    box.append(el(f"w:{side}", **{"w:val": "single", "w:sz": max(2, int(round(size_pt * 8))),
                                  "w:space": space, "w:color": C[color]}))


def rule(container, color="rule", weight=0.7, before=0, after=0):
    p = para(container, before=before, after=after, lh=1)
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    border(p._p.get_or_add_pPr(), "bottom", color, weight)
    return p


def add_link(p, text, url, spec):
    part = p.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = el("w:hyperlink", **{"r:id": r_id})
    run = p.add_run(text)
    style_run(run, spec, color="link")
    h.append(run._r)
    p._p.append(h)


MARK = re.compile(r"(\*\*.+?\*\*|__.+?__|\[[^\]]+\]\([^)]+\)|(?<![*\w])\*[^*]+?\*(?![*\w]))")


def rich(p, text, spec, strong="sans"):
    """Write text with **sans bold**, __serif bold__, *italic*, [link](url)."""
    for part in MARK.split(text):
        if not part:
            continue
        if part.startswith("**"):
            style_run(p.add_run(part[2:-2]), T["strong_sans"] if spec.get("color") != "on_navy_text" else spec,
                      **({} if spec.get("color") != "on_navy_text" else {"font": "sans"}))
        elif part.startswith("__"):
            style_run(p.add_run(part[2:-2]), spec, font="serif_strong")
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            add_link(p, m.group(1), m.group(2), spec)
        elif part.startswith("*"):
            style_run(p.add_run(part[1:-1]), spec, italic=True)
        else:
            style_run(p.add_run(part), spec)
    return p


def cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = el("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(el(f"w:{side}", **{"w:w": tw(v), "w:type": "dxa"}))
    put(tcpr, mar)


def shade(cell, color):
    put(cell._tc.get_or_add_tcPr(), el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": C[color]}))


def table(container, widths, indent=0, borders=None):
    """A fixed-width table with no borders unless asked."""
    t = container.add_table(rows=1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblpr = t._tbl.tblPr
    for old in tblpr.findall(qn("w:tblStyle")):
        tblpr.remove(old)
    put(tblpr, el("w:tblW", **{"w:w": tw(sum(widths)), "w:type": "dxa"}))
    put(tblpr, el("w:tblInd", **{"w:w": tw(indent), "w:type": "dxa"}))
    put(tblpr, el("w:tblLayout", **{"w:type": "fixed"}))
    b = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        spec = (borders or {}).get(side)
        if spec:
            b.append(el(f"w:{side}", **{"w:val": "single", "w:sz": int(spec[1] * 8), "w:space": 0, "w:color": C[spec[0]]}))
        else:
            b.append(el(f"w:{side}", **{"w:val": "nil"}))
    put(tblpr, b)
    cm = el("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        cm.append(el(f"w:{side}", **{"w:w": 0, "w:type": "dxa"}))
    put(tblpr, cm)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(tw(w)))
    for c, w in zip(t.rows[0].cells, widths):
        c.width = Pt(w)
    return t


def add_row(t, widths):
    row = t.add_row()
    for c, w in zip(row.cells, widths):
        c.width = Pt(w)
    return row


def first_para(cell):
    """Cells start with an empty paragraph; reuse it once."""
    p = cell.paragraphs[0]
    if not p.runs and not getattr(cell, "_used", False):
        cell._used = True
        return p
    return None


def cpara(cell, spec=None, **kw):
    p = first_para(cell)
    if p is None:
        return para(cell, spec, **kw)
    pf = p.paragraph_format
    pf.space_before = Pt(kw.get("before", 0))
    pf.space_after = Pt(kw.get("after", 0))
    line = kw.get("lh") or (spec or {}).get("lh")
    if line:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    if kw.get("indent") or kw.get("hanging"):
        pf.left_indent = Pt(kw.get("indent", 0))
        if kw.get("hanging"):
            pf.first_line_indent = Pt(-kw["hanging"])
    if kw.get("right"):
        pf.right_indent = Pt(kw["right"])
    if kw.get("keep"):
        pf.keep_with_next = True
    return p


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.insert(0, el("w:cantSplit"))


# ---------- bullets ----------

def bullet(cell, text, dash_color="accent", spec=None, indent=15.0, after=4.5):
    spec = spec or T["body"]
    p = cpara(cell, spec, indent=indent, hanging=indent, after=after)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(indent))
    style_run(p.add_run("–\t"), {"font": "sans", "size": spec["size"], "color": dash_color})
    rich(p, text, spec)
    return p


# ---------- sections ----------

def spacer(cell):
    """Cells must end in a paragraph; make the one after a nested table take no room."""
    p = cpara(cell)
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style_run(p.add_run(""), {"font": "serif", "size": 1, "color": "ink"})
    return p


def section_frame(doc, label, keep_rows=True, top_rule=True):
    """One labelled section: a rule on top, label in the left gutter, content on the right."""
    sr = T["section_rule"]
    if top_rule:
        sep = para(doc, before=sr["space_above"], after=0, keep=True)
        sep.paragraph_format.line_spacing = Pt(1)
        sep.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        border(sep._p.get_or_add_pPr(), "bottom", sr["color"], sr["weight"])
    widths = [P["label_col"], P["body_col"]]
    t = table(doc, widths)
    row = t.rows[0]
    if keep_rows:
        no_split(row)
    for c in row.cells:
        cell_margins(c, top=sr["space_below"] if top_rule else 0, bottom=0)
    lab = cpara(row.cells[0], before=4)
    style_run(lab.add_run(label), T["section_label"])
    return row.cells[1]


def heading(cell, text, after=8, spec=None):
    p = cpara(cell, after=after, keep=True)
    style_run(p.add_run(text), spec or T["section_heading"])
    return p


def section_lead(cell, text, after=9):
    if text:
        rich(cpara(cell, T["section_lead"], after=after), text, T["section_lead"])


def render_header(doc, jd):
    h = jd["header"]
    top = para(doc, after=0, lh=1)
    top.paragraph_format.line_spacing = Pt(1)
    top.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    border(top._p.get_or_add_pPr(), "top", T["top_bar"]["color"], T["top_bar"]["height"])

    eb = para(doc, before=12, after=14)
    e = T["eyebrow"]
    style_run(eb.add_run(CO["eyebrow_name"]), e, color=e["first_color"])
    for part in (h["function"], h["city"]):
        style_run(eb.add_run(f"  {e['separator']}  "), e, color=e["separator_color"], tracking=0)
        style_run(eb.add_run(part), e)

    t = para(doc, T["title"], after=7, right=P["content_width"] - T["title"]["max_width"], keep=True)
    t.paragraph_format.line_spacing = Pt(T["title"]["lh"])
    t.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style_run(t.add_run(h["title"]), T["title"])

    lead = para(doc, T["lead"], after=18, right=P["content_width"] - T["lead"]["max_width"])
    rich(lead, h["lead"], T["lead"])
    if h.get("punchline"):
        style_run(lead.add_run(" " + h["punchline"]), T["punchline"])

    # meta panel
    n = len(jd["meta"])
    w = P["content_width"] / n
    mt = table(doc, [w] * n, borders={s: ("rule", 0.75) for s in ("top", "left", "bottom", "right")})
    no_split(mt.rows[0])
    for cell, m in zip(mt.rows[0].cells, jd["meta"]):
        shade(cell, "panel")
        cell_margins(cell, top=9, bottom=10, left=10.8, right=6)
        lp = cpara(cell, after=4)
        style_run(lp.add_run(m["label"]), T["meta_label"])
        vp = cpara(cell, T["meta_value"])
        style_run(vp.add_run(m["value"]), T["meta_value"])
    gap = para(doc, after=0)
    gap.paragraph_format.line_spacing = Pt(16)
    gap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def render_about(doc):
    a = CO["about_us"]
    cell = section_frame(doc, a["label"], top_rule=False)
    heading(cell, a["heading"])
    for i, ptxt in enumerate(a["paragraphs"]):
        rich(cpara(cell, T["body"], after=T["body"]["para_gap"] if i < len(a["paragraphs"]) - 1 else 0), ptxt, T["body"])


def render_role(doc, jd):
    s = SKEL["sections"]["the_role"]
    cell = section_frame(doc, s["label"])
    heading(cell, s["heading"])
    section_lead(cell, jd["the_role"]["lead"])
    ps = jd["the_role"]["paragraphs"]
    for i, ptxt in enumerate(ps):
        rich(cpara(cell, T["body"], after=T["body"]["para_gap"] if i < len(ps) - 1 else 0), ptxt, T["body"])


def render_level(doc, jd):
    s = SKEL["sections"]["the_level"]
    lv = jd["the_level"]
    cell = section_frame(doc, s["label"])
    heading(cell, s["heading"].format(level_name=lv["level_name"]))
    section_lead(cell, lv["lead"], after=10)
    k = T["level_key"]
    widths = [k["col"], P["body_col"] - k["col"]]
    t = table(cell, widths, borders={"top": ("rule", 0.75), "bottom": ("rule", 0.75), "insideH": ("rule", 0.75)})
    for i, r in enumerate(lv["rows"]):
        row = t.rows[0] if i == 0 else add_row(t, widths)
        no_split(row)
        row.height = Pt(k["row_h"])
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(c, top=6.5, bottom=6.5)
        style_run(cpara(row.cells[0]).add_run(r["key"]), k)
        rich(cpara(row.cells[1], T["body"]), r["text"], T["body"])
    spacer(cell)


def render_work(doc, jd):
    s = SKEL["sections"]["your_work"]
    w = jd["your_work"]
    cell = section_frame(doc, s["label"], keep_rows=False)   # the one section long enough to need a page break
    heading(cell, w["heading"], after=10)
    for g in w["groups"]:
        p = cpara(cell, before=0, after=8, keep=True)
        border(p._p.get_or_add_pPr(), "top", "rule", 0.7, space=10)
        p.paragraph_format.space_before = Pt(4)
        style_run(p.add_run(g["label"]), T["group_label"])
        for i, b in enumerate(g["bullets"]):
            bullet(cell, b, after=T["bullet"]["gap"] if i < len(g["bullets"]) - 1 else 10)


def render_bar(doc, jd):
    s = SKEL["sections"]["the_bar"]
    bar = jd["the_bar"]
    cell = section_frame(doc, s["label"])
    box = T["bar_box"]
    t = table(cell, [P["body_col"]])
    inner = t.rows[0].cells[0]
    shade(inner, box["fill"])
    cell_margins(inner, top=box["padding"], bottom=box["padding"], left=box["padding"], right=box["padding"])
    style_run(cpara(inner, after=9).add_run(s["fixed"]["eyebrow"]), T["bar_eyebrow"])
    style_run(cpara(inner, after=7, keep=True).add_run(s["fixed"]["heading"]), T["bar_heading"])
    rich(cpara(inner, T["bar_text"], after=12), bar["intro"], T["bar_text"])
    for i, tr in enumerate(bar["traits"]):
        p = cpara(inner, before=0, after=4, keep=True)
        border(p._p.get_or_add_pPr(), "top", box["rule"], 0.7, space=10)
        p.paragraph_format.space_before = Pt(2)
        style_run(p.add_run(tr["title"]), T["bar_trait"])
        rich(cpara(inner, T["bar_text"], after=12 if i < len(bar["traits"]) - 1 else 0), tr["text"], T["bar_text"])
    spacer(cell)


def render_ai(doc, jd):
    s = SKEL["sections"]["ai"]
    cell = section_frame(doc, s["label"])
    heading(cell, s["heading"])
    section_lead(cell, jd["ai"]["lead"], after=8)
    bs = jd["ai"]["bullets"]
    for i, b in enumerate(bs):
        bullet(cell, b, after=T["bullet"]["gap"] if i < len(bs) - 1 else 0)


def render_minimum(doc, jd):
    s = SKEL["sections"]["minimum"]
    m = jd["minimum"]
    cell = section_frame(doc, s["label"])
    heading(cell, s["heading"])
    section_lead(cell, m.get("note"), after=10)
    items = m["items"]
    half = (len(items) + 1) // 2
    widths = [215.0, P["body_col"] - 215.0]
    t = table(cell, widths)
    left, right = t.rows[0].cells
    cell_margins(left, right=14)
    for col, chunk in ((left, items[:half]), (right, items[half:])):
        for i, it in enumerate(chunk):
            bullet(col, it, after=T["bullet"]["gap"] + 1 if i < len(chunk) - 1 else 0)
    spacer(cell)


def render_preferred(doc, jd):
    s = SKEL["sections"]["preferred"]
    cell = section_frame(doc, s["label"])
    heading(cell, s["heading"])
    items = jd["preferred"]["items"]
    for i, it in enumerate(items):
        bullet(cell, it, after=T["bullet"]["gap"] if i < len(items) - 1 else 0)


def render_toolkit(doc, jd):
    s = SKEL["sections"]["toolkit"]
    tk = jd["toolkit"]
    cell = section_frame(doc, tk.get("label", "OUR STACK"))
    heading(cell, s["heading"], after=6)
    section_lead(cell, tk.get("note"), after=10)
    k = T["stack_key"]
    widths = [k["col"], P["body_col"] - k["col"]]
    t = table(cell, widths, borders={"insideH": ("rule", 0.75), "bottom": ("rule", 0.75)})
    for i, r in enumerate(tk["rows"]):
        row = t.rows[0] if i == 0 else add_row(t, widths)
        no_split(row)
        row.height = Pt(k["row_h"])
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(c, top=6, bottom=6)
        style_run(cpara(row.cells[0]).add_run(r["key"]), k)
        style_run(cpara(row.cells[1]).add_run(r["value"]), T["stack_value"])
    spacer(cell)


def render_assessment(doc, jd):
    s = SKEL["sections"]["assessment"]
    a = jd["assessment"]
    cell = section_frame(doc, s["label"])
    heading(cell, s["heading"])
    section_lead(cell, a["lead"], after=12)
    box = T["card_box"]
    cols = box["cols"]
    w = P["body_col"] / cols
    widths = [w] * cols
    b = ("rule", 0.75)
    t = table(cell, widths, borders={s_: b for s_ in ("top", "left", "bottom", "right", "insideH", "insideV")})
    cards = a["cards"]
    for r in range(0, len(cards), cols):
        row = t.rows[0] if r == 0 else add_row(t, widths)
        no_split(row)
        for c, card in zip(row.cells, cards[r:r + cols]):
            cell_margins(c, top=14, bottom=14, left=box["padding"] + 1.5, right=box["padding"])
            style_run(cpara(c, after=6).add_run(card["title"]), T["card_title"])
            rich(cpara(c, T["card_text"]), card["text"], T["card_text"])
    spacer(cell)


def render_not_for_you(doc, jd):
    s = SKEL["sections"]["not_for_you"]
    cell = section_frame(doc, s["label"])
    box = T["panel_box"]
    t = table(cell, [P["body_col"]], borders={x: ("rule", 0.75) for x in ("top", "left", "bottom", "right")})
    inner = t.rows[0].cells[0]
    shade(inner, box["fill"])
    cell_margins(inner, top=box["padding"], bottom=box["padding"] + 4, left=box["padding"] + 2, right=box["padding"])
    style_run(cpara(inner, after=9, keep=True).add_run(s["heading"]), T["panel_title"])
    items = jd["not_for_you"]["items"]
    for i, it in enumerate(items):
        bullet(inner, it, dash_color=T["panel_text"]["dash_color"], spec=T["panel_text"],
               after=T["bullet"]["gap"] if i < len(items) - 1 else 0)
    spacer(cell)


def render_eeo_and_footer(doc, jd):
    """Equal opportunity, then the footer line, in one table so the footer never sits alone on a page."""
    e = CO["eeo"]
    cell = section_frame(doc, e["label"])
    rich(cpara(cell, T["eeo"], keep=True), e["text"], T["eeo"])
    tbl = cell._tc.getparent().getparent()
    from docx.table import Table
    table_obj = Table(tbl, cell._parent)
    widths = [P["label_col"], P["body_col"]]
    row = add_row(table_obj, widths)
    no_split(row)
    merged = row.cells[0].merge(row.cells[1])
    sr = T["section_rule"]
    border(merged._tc.get_or_add_tcPr(), "top", sr["color"], sr["weight"])
    cell_margins(merged, top=sr["space_below"], bottom=0)
    # keep the two rows on one page
    for p_ in cell.paragraphs:
        p_.paragraph_format.keep_with_next = True
    f = T["footer"]
    gap = "\u2003\u2003"
    fp = cpara(merged)
    style_run(fp.add_run(CO["footer"]["left"] + gap + jd["footer_tag"] + gap + CO["footer"]["right"]), f)
    # first row's cells also get the eeo bottom gap
    for c in table_obj.rows[0].cells:
        cell_margins(c, top=sr["space_below"], bottom=sr["space_above"])


# ---------- document setup ----------

def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Pt(P["width"]), Pt(P["height"])
    s.top_margin, s.bottom_margin = Pt(P["margin_top"]), Pt(P["margin_bottom"])
    s.left_margin, s.right_margin = Pt(P["margin_left"]), Pt(P["margin_right"])
    s.header_distance = s.footer_distance = Pt(20)
    normal = doc.styles["Normal"]
    normal.font.name = F["serif"]
    normal.font.size = Pt(T["body"]["size"])
    normal.font.color.rgb = RGBColor.from_string(C["ink"])
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), F["serif"])
    normal.paragraph_format.space_after = Pt(0)
    # British English for the spell checker in Word
    put(rpr, el("w:lang", **{"w:val": "en-GB"}))
    # drop the template's empty first paragraph
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)


def embed_fonts(doc):
    """Embed the brand TTFs (obfuscated, per ECMA-376 17.8.1) so Word shows them without installing."""
    font_dir = ROOT / "brand/fonts"
    ft_part = doc.part.part_related_by(RT.FONT_TABLE)
    ft = parse_xml(ft_part.blob)
    used = set()
    for r in doc.element.body.iter(qn("w:rFonts")):
        used.add(r.get(qn("w:ascii")))
    used.add(F["serif"])
    for name in sorted(used):
        files = TOK["font_files"].get(name)
        if not files:
            continue
        font_el = None
        for f in ft.findall(qn("w:font")):
            if f.get(qn("w:name")) == name:
                font_el = f
        if font_el is None:
            font_el = el("w:font", **{"w:name": name})
            ft.append(font_el)
        for kind, fname in files.items():
            key = "{" + str(uuid.uuid4()).upper() + "}"
            data = bytearray((font_dir / fname).read_bytes())
            kb = bytes.fromhex(key.strip("{}").replace("-", ""))[::-1]
            for i in range(32):
                data[i] ^= kb[i % 16]
            idx = len(ft_part.rels) + 1
            part = Part(PackURI(f"/word/fonts/font{idx}.odttf"),
                        "application/vnd.openxmlformats-officedocument.obfuscatedFont", bytes(data), ft_part.package)
            rid = ft_part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
            tag = {"regular": "w:embedRegular", "bold": "w:embedBold", "italic": "w:embedItalic"}[kind]
            font_el.append(el(tag, **{"r:id": rid, "w:fontKey": key}))
    ft_part._blob = etree.tostring(ft, xml_declaration=True, encoding="UTF-8", standalone=True)
    put(doc.settings.element, el("w:embedTrueTypeFonts"))


def render(jd, out):
    doc = Document()
    setup(doc)
    render_header(doc, jd)
    render_about(doc)
    render_role(doc, jd)
    render_level(doc, jd)
    render_work(doc, jd)
    render_bar(doc, jd)
    render_ai(doc, jd)
    render_minimum(doc, jd)
    render_preferred(doc, jd)
    render_toolkit(doc, jd)
    render_assessment(doc, jd)
    render_not_for_you(doc, jd)
    render_eeo_and_footer(doc, jd)
    embed_fonts(doc)
    cp = doc.core_properties
    cp.title = f"JD - {CO['short_name']} {jd['header']['title']}"
    cp.author = CO["legal_name"]
    cp.subject = "Job description"
    doc.save(out)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    render(json.loads(Path(sys.argv[1]).read_text()), sys.argv[2])
    print(f"wrote {sys.argv[2]}")
