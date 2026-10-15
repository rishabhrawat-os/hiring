#!/usr/bin/env python3
"""Collect the JD and every CV into a screening run folder.

Accepts any mix of files, folders and .zip archives (PDF, DOCX, TXT, MD).
Each CV gets an ID (C01, C02, ...) and its own folder with the extracted
text. The ID-to-file map stays in manifest.json, in the run folder only.

    python3 ingest.py <run> --jd <jd file> <cv paths...>
"""
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

CV_TYPES = {".pdf", ".docx", ".txt", ".md"}


def pdf_text(path):
    try:
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(str(path)).pages)
    except ImportError:
        pass
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True)
        return r.stdout
    return ""


def docx_text(path):
    try:
        import docx
    except ImportError:
        return ""
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def to_text(path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return pdf_text(path)
    if ext == ".docx":
        return docx_text(path)
    return path.read_text(errors="replace")


def expand(paths, scratch):
    out = []
    for p in map(Path, paths):
        if p.is_dir():
            out += sorted(f for f in p.rglob("*") if f.is_file())
        elif p.suffix.lower() == ".zip":
            dest = scratch / p.stem
            with zipfile.ZipFile(p) as z:
                for m in z.namelist():
                    target = (dest / m).resolve()
                    if dest.resolve() not in target.parents and target != dest.resolve():
                        continue          # skip unsafe paths inside the archive
                    if m.endswith("/"):
                        continue
                    target.parent.mkdir(parents=True, exist_ok=True)
                    target.write_bytes(z.read(m))
            out += sorted(f for f in dest.rglob("*") if f.is_file())
        elif p.is_file():
            out.append(p)
        else:
            print(f"WARN not found: {p}")
    return [f for f in out if f.suffix.lower() in CV_TYPES and not f.name.startswith(("~$", "."))]


def main():
    a = sys.argv[1:]
    if len(a) < 4 or "--jd" not in a:
        sys.exit(__doc__)
    run = Path(a[0])
    jd = Path(a[a.index("--jd") + 1])
    cvs = [x for i, x in enumerate(a[1:], 1) if x != "--jd" and a[i - 1] != "--jd"]
    run.mkdir(parents=True, exist_ok=True)
    (run / "input").mkdir(exist_ok=True)

    jd_text = to_text(jd)
    shutil.copy(jd, run / "input" / jd.name)
    (run / "00-jd.txt").write_text(jd_text)

    manifest = json.loads((run / "manifest.json").read_text()) if (run / "manifest.json").exists() else {"candidates": []}
    seen = {c["source"] for c in manifest["candidates"]}
    files = expand(cvs, run / "input" / "unzipped")
    n = len(manifest["candidates"])
    for f in files:
        if str(f.resolve()) in seen:
            continue
        n += 1
        cid = f"C{n:02d}"
        d = run / "candidates" / cid
        d.mkdir(parents=True, exist_ok=True)
        text = to_text(f)
        (d / "cv.txt").write_text(text)
        shutil.copy(f, d / f"original{f.suffix.lower()}")
        status = "ingested" if len(text.split()) >= 50 else "needs_manual_read"
        (d / "status.json").write_text(json.dumps({"id": cid, "step": "ingested", "status": status}, indent=1))
        manifest["candidates"].append({"id": cid, "source": str(f.resolve()), "label": f.stem, "words": len(text.split()), "status": status})
        print(f"{cid}  {f.name}  {len(text.split())} words  {status}")
    (run / "manifest.json").write_text(json.dumps(manifest, indent=1))
    print(json.dumps({"run": str(run), "jd_words": len(jd_text.split()), "candidates": len(manifest["candidates"])}))


if __name__ == "__main__":
    main()
